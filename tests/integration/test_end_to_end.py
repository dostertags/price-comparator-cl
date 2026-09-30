"""Extremo a extremo SIN red: Excel -> CLI -> scrapers reales (parsers) sobre fixtures -> Excel."""

import io

from openpyxl import Workbook, load_workbook

from price_comparator import cli
from price_comparator.excel.writer import RESULT_COLUMNS
from price_comparator.scrapers.falabella import FalabellaScraper
from price_comparator.scrapers.hites import HitesScraper
from price_comparator.scrapers.sodimac import SodimacScraper
from tests.fakes import FakeFetcher, ok


def test_excel_to_excel_with_real_parsers_and_recorded_pages(tmp_path, fx):
    def factory(_settings):
        return [
            SodimacScraper(FakeFetcher({"sodimac.cl": ok(fx("sodimac_results.html"))})),
            FalabellaScraper(FakeFetcher({"falabella.com": ok(fx("falabella_results.html"))})),
            HitesScraper(FakeFetcher({"hites.com": ok(fx("hites_results.html"))})),
        ]

    src = tmp_path / "lista.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Producto", "Cantidad"])
    ws.append(["Taladro percutor eléctrico 13 mm 750W", 2])
    ws.append(["Taladro percutor eléctrico 13 mm 750W", 3])  # duplicado
    ws.append(["Producto que no existe zzqxkjwv", 1])
    wb.save(src)

    out, err = io.StringIO(), io.StringIO()
    code = cli.main([str(src)], environ={}, scrapers_factory=factory, stdout=out, stderr=err)

    assert code == 0, err.getvalue()
    produced = next(tmp_path.glob("lista_resultado_*.xlsx"))
    sheet = load_workbook(produced)["Resultados"]
    col = {name: i for i, name in enumerate(RESULT_COLUMNS, start=1)}

    row2 = {name: sheet.cell(2, i).value for name, i in col.items()}
    assert row2["Estado"] == "Encontrado"
    assert row2["Tienda ganadora"] == "Sodimac"
    assert row2["Mejor precio (CLP)"] == 51_990  # el Flowmak de Hites es más barato pero sin 750W
    assert row2["Total (CLP)"] == 2 * 51_990
    assert sheet.cell(3, col["Mejor precio (CLP)"]).value == 51_990  # duplicado, misma respuesta
    assert sheet.cell(3, col["Total (CLP)"]).value == 3 * 51_990

    row4 = {name: sheet.cell(4, i).value for name, i in col.items()}
    assert row4["Mejor precio (CLP)"] is None
    assert row4["Estado"] in ("Baja confianza", "Sin resultados")

    assert load_workbook(produced).sheetnames == ["Resultados", "Ofertas", "Resumen"]
    assert src.exists() and load_workbook(src).sheetnames == ["Sheet"]  # entrada intacta
