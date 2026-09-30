import io
import logging
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from price_comparator import cli
from price_comparator.logging_setup import configure_logging
from price_comparator.models import ScrapeResult, ScrapeStatus
from tests.factories import make_offer
from tests.fakes import FakeScraper, result

TALADRO = "Taladro percutor 750W"


@pytest.fixture(autouse=True)
def _clean_logging():
    yield
    logger = logging.getLogger("price_comparator")
    for h in [h for h in logger.handlers if h.get_name() == "price_comparator_cli"]:
        logger.removeHandler(h)
        h.close()


def make_input(path, rows=(TALADRO, "Sierra circular 1200W"), header="Producto"):
    wb = Workbook()
    ws = wb.active
    ws.append([header])
    for r in rows:
        ws.append([r])
    wb.save(path)
    return path


def good_scrapers(_settings):
    def behaviour(query):
        return result("Sodimac", make_offer("Sodimac", query, 50_000))

    return [FakeScraper("Sodimac", behaviour)]


def broken_scrapers(_settings):
    return [FakeScraper("Sodimac", ScrapeResult("Sodimac", ScrapeStatus.BLOCKED, detail="x"))]


def run(argv, scrapers=good_scrapers, env=None):
    out, err = io.StringIO(), io.StringIO()
    code = cli.main(argv, environ=env or {}, scrapers_factory=scrapers, stdout=out, stderr=err)
    return code, out.getvalue(), err.getvalue()


def test_success_writes_report_next_to_input_and_exits_zero(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    code, out, _ = run([str(src)])
    assert code == 0
    produced = list(tmp_path.glob("lista_resultado_*.xlsx"))
    assert len(produced) == 1
    assert str(produced[0]) in out
    ws = load_workbook(produced[0])["Resultados"]
    assert ws.max_row == 3
    assert ws.cell(2, 3).value == 50_000


def test_custom_output_path(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    dest = tmp_path / "mio.xlsx"
    code, _, _ = run([str(src), "--output", str(dest)])
    assert code == 0 and dest.exists()


def test_missing_column_is_a_usage_error(tmp_path):
    src = make_input(tmp_path / "lista.xlsx", header="Código")
    code, _, err = run([str(src)])
    assert code == 2
    assert "Código" in err


def test_column_option_is_honoured(tmp_path):
    src = make_input(tmp_path / "lista.xlsx", header="Detalle")
    assert run([str(src), "--column", "Detalle"])[0] == 0


def test_missing_file_is_a_usage_error(tmp_path):
    code, _, err = run([str(tmp_path / "nada.xlsx")])
    assert code == 2 and "no existe" in err


def test_invalid_configuration_is_a_usage_error(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    code, _, err = run([str(src)], env={"PC_WORKERS": "99"})
    assert code == 2 and "PC_WORKERS" in err


def test_cli_flags_override_environment(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    seen = {}

    def factory(settings):
        seen["min_score"] = settings.min_score
        seen["workers"] = settings.workers
        return good_scrapers(settings)

    run(
        [str(src), "--min-score", "0.9", "--workers", "2"],
        scrapers=factory,
        env={"PC_MIN_SCORE": "0.5"},
    )
    assert seen == {"min_score": 0.9, "workers": 2}


def test_when_every_store_is_down_it_fails_loudly_and_writes_no_excel(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    code, _, err = run([str(src)], scrapers=broken_scrapers)
    assert code == 1
    assert "ninguna tienda" in err.lower()
    assert list(tmp_path.glob("*_resultado_*.xlsx")) == []


def test_progress_goes_to_stderr_and_path_to_stdout(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    _, out, err = run([str(src)])
    assert "[1/2]" in err or "[2/2]" in err
    assert "[1/2]" not in out


def test_unexpected_errors_are_reported_without_a_traceback_and_exit_one(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")

    def boom(_settings):
        raise RuntimeError("kaput")

    code, _, err = run([str(src)], scrapers=boom)
    assert code == 1
    assert "Traceback" not in err and "--verbose" in err


def test_help_exits_zero_and_bad_flag_exits_two():
    with pytest.raises(SystemExit) as ok:
        cli.main(["--help"], environ={})
    assert ok.value.code == 0
    with pytest.raises(SystemExit) as bad:
        cli.main(["x.xlsx", "--nope"], environ={})
    assert bad.value.code == 2


def test_importing_cli_does_not_configure_logging():
    """En un intérprete limpio: importar el paquete no agrega ningún handler."""
    code = (
        "import logging, price_comparator.cli;"
        "print(len(logging.getLogger('price_comparator').handlers),"
        " len(logging.getLogger().handlers))"
    )
    src = Path(__file__).resolve().parents[2] / "src"
    env = {**os.environ, "PYTHONPATH": str(src)}
    proc = subprocess.run(
        [sys.executable, "-c", code], capture_output=True, text=True, env=env, check=True
    )
    assert proc.stdout.split() == ["0", "0"]


def test_configure_logging_is_idempotent_and_leaves_root_alone():
    root_before = list(logging.getLogger().handlers)
    configure_logging("DEBUG")
    configure_logging("INFO")
    logger = logging.getLogger("price_comparator")
    cli_handlers = [h for h in logger.handlers if h.get_name() == "price_comparator_cli"]
    assert len(cli_handlers) == 1
    assert logger.level == logging.INFO
    assert logging.getLogger().handlers == root_before
    for h in cli_handlers:
        logger.removeHandler(h)


def test_unwritable_log_file_is_a_usage_error_not_a_traceback(tmp_path):
    src = make_input(tmp_path / "lista.xlsx")
    code, _, err = run([str(src), "--log-file", str(tmp_path / "no" / "existe" / "x.log")])
    assert code == 2
    assert "log" in err.lower() and "Traceback" not in err


@pytest.mark.parametrize("value", ["0", "-5"])
def test_non_positive_max_rows_is_a_usage_error(tmp_path, value):
    src = make_input(tmp_path / "lista.xlsx")
    code, _, err = run([str(src), "--max-rows", value])
    assert code == 2 and "max-rows" in err
