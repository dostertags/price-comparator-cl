"""Lectura del Excel de entrada."""

from __future__ import annotations

import logging
import math
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from price_comparator.comparison.text import normalize_name, strip_accents

log = logging.getLogger("price_comparator")

MAX_ROWS_HARD_CAP = 2000
DEFAULT_MAX_ROWS = 500
DEFAULT_MAX_BYTES = 5_000_000
NAME_MAX_LEN = 200
MAX_QUANTITY = 10_000_000
MAX_UNCOMPRESSED_BYTES = 100_000_000
MAX_ZIP_ENTRIES = 1000

_PRODUCT_ALIASES = ("producto", "nombre", "descripcion", "product", "name")
_QUANTITY_ALIASES = ("cantidad", "cant", "qty", "quantity")
_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")  # .xls o .xlsx cifrado


class InputError(Exception):
    """El archivo de entrada no se puede usar; el mensaje está pensado para el usuario."""


@dataclass(frozen=True)
class ProductRow:
    """Una fila de entrada."""

    row_number: int
    original: str
    name: str
    quantity: float | None = None
    notes: tuple[str, ...] = ()


def _key(value: Any) -> str:
    return strip_accents(str(value)).strip().lower() if value is not None else ""


def _quantity(value: Any) -> tuple[float | None, str]:
    """Interpreta una celda de cantidad: `(cantidad, nota)`; la nota explica un descarte."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, ""
    invalid = (None, "cantidad no válida (se ignora)")
    if isinstance(value, bool):
        return invalid
    try:
        number = (
            float(str(value).strip().replace(",", ".")) if isinstance(value, str) else float(value)
        )
    except (ValueError, OverflowError):
        return invalid
    if not math.isfinite(number) or not 0 < number <= MAX_QUANTITY:
        return invalid
    return number, ""


def _check_file(path: Path, max_bytes: int) -> None:
    if path.name.startswith("~$"):
        raise InputError(
            f"«{path.name}» es un archivo temporal que Excel crea mientras el original está "
            "abierto; usa el archivo real."
        )
    if path.suffix.lower() != ".xlsx":
        raise InputError(
            f"Solo se soportan archivos .xlsx (recibí «{path.suffix or 'sin extensión'}»)."
        )
    if not path.is_file():
        raise InputError(f"El archivo no existe: {path}")
    size = path.stat().st_size
    if size > max_bytes:
        raise InputError(
            f"El archivo pesa {size / 1e6:.1f} MB y el máximo es {max_bytes / 1e6:.1f} MB."
        )
    try:
        with path.open("rb") as fh:
            head = fh.read(8)
    except PermissionError as exc:
        raise InputError(f"Sin permiso para leer {path}") from exc
    if head == _OLE_MAGIC:
        raise InputError(
            "El archivo parece estar protegido con contraseña (o es un .xls renombrado); "
            "quítale la contraseña y guárdalo como .xlsx."
        )


def read_products(
    path: Path | str,
    *,
    column: str | None = None,
    sheet: str | None = None,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_bytes: int = DEFAULT_MAX_BYTES,
    max_uncompressed_bytes: int = MAX_UNCOMPRESSED_BYTES,
) -> list[ProductRow]:
    """Lee la lista de productos de un `.xlsx`.

    Args:
        path: archivo de entrada.
        column: nombre de la columna con el producto; si falta, se busca entre los alias
            habituales (Producto, Nombre, Descripción, Product, Name).
        sheet: hoja a leer; por defecto la primera.
        max_rows: máximo de filas de datos (tope duro: 2000).
        max_bytes: tamaño máximo del archivo.
        max_uncompressed_bytes: tamaño máximo descomprimido (defensa contra bombas zip).

    Raises:
        InputError: con un mensaje claro para cualquier problema del archivo.
    """
    path = Path(path)
    if max_rows > MAX_ROWS_HARD_CAP:
        raise InputError(f"--max-rows no puede superar {MAX_ROWS_HARD_CAP}.")
    _check_file(path, max_bytes)

    try:
        _check_zip(path, max_uncompressed_bytes)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is not None and sheet not in wb.sheetnames:
                raise InputError(
                    f"No existe la hoja «{sheet}». Hojas disponibles: {', '.join(wb.sheetnames)}"
                )
            ws: Any = wb[sheet] if sheet is not None else wb.worksheets[0]
            return _read_sheet(ws, column, max_rows)
        finally:
            wb.close()
    except InputError:
        raise
    except Exception as exc:  # noqa: BLE001 - frontera con un archivo no confiable
        # Zip roto, XML mal formado, entidades prohibidas por defusedxml, etc.: nunca un traceback.
        log.debug("no se pudo leer el xlsx", exc_info=True)
        raise InputError(
            "El archivo no parece un .xlsx válido. Ábrelo en Excel y guárdalo de nuevo como .xlsx."
        ) from exc


def _check_zip(path: Path, max_uncompressed_bytes: int) -> None:
    """Rechaza «bombas zip»: poco en disco, enorme al descomprimir."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
    if len(entries) > MAX_ZIP_ENTRIES:
        raise InputError(f"El archivo trae demasiados archivos internos ({len(entries)}).")
    if sum(e.file_size for e in entries) > max_uncompressed_bytes:
        raise InputError(
            f"El archivo ocupa más de {max_uncompressed_bytes / 1e6:.0f} MB descomprimido."
        )


def _read_sheet(ws: Any, column: str | None, max_rows: int) -> list[ProductRow]:
    rows = enumerate(ws.iter_rows(values_only=True), start=1)
    header: tuple[Any, ...] | None = None
    for _, cells in rows:
        if any(c not in (None, "") for c in cells):
            header = cells
            break
    if header is None:
        raise InputError("La hoja está vacía.")

    keys = [_key(h) for h in header]

    def find(aliases: tuple[str, ...]) -> int | None:
        return next((i for i, k in enumerate(keys) if k in aliases), None)

    wanted = (_key(column),) if column else _PRODUCT_ALIASES
    name_idx = find(wanted)
    if name_idx is None:
        found = ", ".join(str(h) for h in header if h not in (None, "")) or "(ninguna)"
        accepted = column or ", ".join(a.capitalize() for a in _PRODUCT_ALIASES)
        raise InputError(
            f"No encontré la columna del producto. Columnas encontradas: {found}. "
            f"Se acepta: {accepted} (o indica otra con --column)."
        )
    qty_idx = find(_QUANTITY_ALIASES)

    def cell(cells: tuple[Any, ...], idx: int | None) -> Any:
        return cells[idx] if idx is not None and idx < len(cells) else None

    products: list[ProductRow] = []
    for row_number, cells in rows:
        if not any(c not in (None, "") for c in cells):
            continue
        if len(products) >= max_rows:
            raise InputError(
                f"El archivo tiene más de {max_rows} productos. Divídelo o sube --max-rows "
                f"(tope {MAX_ROWS_HARD_CAP})."
            )
        raw = cell(cells, name_idx)
        original = "" if raw is None else str(raw)
        notes: list[str] = []
        if len(normalize_name(raw, max_len=10**6)) > NAME_MAX_LEN:
            notes.append(f"nombre recortado a {NAME_MAX_LEN} caracteres")
        quantity, quantity_note = _quantity(cell(cells, qty_idx))
        if quantity_note:
            notes.append(quantity_note)
        products.append(
            ProductRow(
                row_number=row_number,
                original=original,
                name=normalize_name(raw, NAME_MAX_LEN),
                quantity=quantity,
                notes=tuple(notes),
            )
        )
    if not products:
        raise InputError("La hoja solo tiene el encabezado: no hay productos.")
    return products
